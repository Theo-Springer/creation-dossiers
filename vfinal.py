from pathlib import Path
import json
import os
import subprocess
import sys
import tkinter as tk
import time
import threading
from tkinter import ttk, messagebox


# ============================================================
# Réglages : lus depuis config.json
# ============================================================

# Le dossier du programme : à côté du .exe une fois compilé, à côté de main.py sinon
if getattr(sys, "frozen", False):
    DOSSIER_APP = Path(sys.executable).parent
else:
    DOSSIER_APP = Path(__file__).parent

INTERDITS = '\\/:*?"<>|'
version = "1.3.0"


def erreur_fatale(titre, message):
    """Affiche une erreur puis ferme le programme (utilisée avant que la fenêtre existe)."""
    racine = tk.Tk()
    racine.withdraw()       # cache la fenêtre vide que tkinter crée pour la boîte de dialogue
    messagebox.showerror(titre, message)
    sys.exit()


chemin_config = DOSSIER_APP / "config.json"
try:
    with open(chemin_config, encoding="utf-8") as fichier:
        config = json.load(fichier)
    BASE = DOSSIER_APP / config["chemin_devis"]     # le dossier où se trouvent les dossiers clients
    SOUS_DOSSIERS = config["sous_dossiers"]
    EXCEL = config.get("excel")     # facultatif : None si la section "excel" n'existe pas
except FileNotFoundError:
    erreur_fatale("Configuration introuvable", f"Le fichier {chemin_config} est introuvable.")
except json.JSONDecodeError as erreur:
    erreur_fatale("Configuration invalide", f"config.json est mal écrit (ligne {erreur.lineno}).\n\nVérifie les virgules et les guillemets.")
except KeyError as erreur:
    erreur_fatale("Configuration incomplète", f"Il manque la clé {erreur} dans config.json.")


# ============================================================
# Logique : dossiers
# ============================================================

def nettoyer(texte):
    """Met en majuscules et ne garde qu'un espace entre les mots."""
    return " ".join(texte.upper().split())


def liste_clients():
    """Renvoie les noms des dossiers clients, triés par ordre alphabétique."""
    clients = []
    for dossier in BASE.iterdir():
        if dossier.is_dir():
            clients.append(dossier.name)
    return sorted(clients, key=str.upper)   # key=str.upper : trie sans tenir compte des majuscules


def nom_interdit(texte):
    """True si le texte ne peut pas servir de nom de dossier Windows."""
    for lettre in texte:
        if lettre in INTERDITS:
            return True
    return texte.endswith(".")


def affaire_existe(numero):
    """Renvoie le dossier qui porte déjà ce numéro (chez n'importe quel client), ou None."""
    for dossier in BASE.glob(f"*/{numero}*"):
        if dossier.is_dir():
            return dossier
    return None


def numeros_existants():
    """Renvoie l'ensemble des numéros d'affaire déjà présents, tous clients confondus."""
    numeros = set()
    for client in BASE.iterdir():
        if client.is_dir():
            for dossier in client.iterdir():
                if dossier.is_dir():
                    numeros.add(dossier.name[:5])
    return numeros


def verifier(client, numero, initiales, objet):
    """Renvoie la liste des erreurs. Liste vide = tout est bon."""
    erreurs = []
    numero_valide = len(numero) == 5 and numero.isdigit()

    if client == "":
        erreurs.append("Le client est vide")
    elif nom_interdit(client):
        erreurs.append(f"Le nom du client contient un caractère interdit ({INTERDITS}) ou finit par un point")

    if not numero_valide:
        erreurs.append("Le numéro d'affaire doit contenir 5 chiffres")

    if not initiales.isalpha():
        erreurs.append("Les initiales doivent contenir uniquement des lettres")

    if objet == "":
        erreurs.append("L'objet est vide")
    elif nom_interdit(objet):
        erreurs.append(f"L'objet contient un caractère interdit ({INTERDITS}) ou finit par un point")

    if numero_valide:
        existant = affaire_existe(numero)
        if existant is not None:
            erreurs.append(f"Le numéro d'affaire existe déjà chez {existant.parent.name} : {existant.name}")

    return erreurs


def creer_affaire(client, numero, initiales, objet):
    """Crée le dossier d'affaire et ses sous-dossiers. Renvoie son chemin."""
    dossier_affaire = BASE / client / f"{numero}{initiales} {objet}"
    dossier_affaire.mkdir(parents=True)     # parents=True : crée aussi le client s'il est nouveau
    for sous_dossier in SOUS_DOSSIERS:
        (dossier_affaire / sous_dossier).mkdir()
    return dossier_affaire


def ouvrir_dossier(chemin):
    """Ouvre le dossier dans l'Explorateur (Windows) ou le Finder (Mac)."""
    if sys.platform == "win32":
        os.startfile(chemin)
    else:
        subprocess.run(["open", str(chemin)])


# ============================================================
# Logique : Excel
# ============================================================

def texte_case(ligne, index):
    """Renvoie le contenu d'une case sous forme de texte, ou "" si elle est vide."""
    if index >= len(ligne) or ligne[index] is None:
        return ""
    return str(ligne[index]).strip()


def lire_excel():
    """Lit l'Excel et renvoie les lignes complètes : (numéro de ligne, client, affaire, objet)."""
    # Importé ici et pas en haut du fichier : openpyxl n'est chargé que si on lit vraiment l'Excel
    from openpyxl import load_workbook
    from openpyxl.utils import column_index_from_string

    # "A" -> 0, "B" -> 1... : la position de la colonne dans le tuple renvoyé par openpyxl
    col_client = column_index_from_string(EXCEL["colonne_client"]) - 1
    col_affaire = column_index_from_string(EXCEL["colonne_affaire"]) - 1
    col_objet = column_index_from_string(EXCEL["colonne_objet"]) - 1
    derniere_colonne = max(col_client, col_affaire, col_objet) + 1
    premiere_ligne = EXCEL["premiere_ligne"]

    classeur = load_workbook(DOSSIER_APP / EXCEL["chemin"], read_only=True, data_only=True)
    try:
        if EXCEL["feuille"]:
            feuille = classeur[EXCEL["feuille"]]
        else:
            feuille = classeur.active       # pas de feuille précisée : celle ouverte par défaut

        lignes = []
        toutes_les_lignes = feuille.iter_rows(min_row=premiere_ligne, max_col=derniere_colonne, values_only=True)
        for numero_ligne, ligne in enumerate(toutes_les_lignes, start=premiere_ligne):
            client = texte_case(ligne, col_client)
            affaire = texte_case(ligne, col_affaire)
            objet = texte_case(ligne, col_objet)
            if client and affaire and objet:        # ligne incomplète : on l'ignore
                lignes.append((numero_ligne, client, affaire, objet))
    finally:
        classeur.close()        # finally : exécuté même si une erreur arrive pendant la lecture

    return lignes


def comparer_excel():
    """Trie les lignes de l'Excel en deux listes : à créer, à corriger. Les affaires existantes sont ignorées."""
    existants = numeros_existants()

    # Pour retrouver le vrai nom du dossier quel que soit l'écriture dans l'Excel :
    # "DUPONT INDUSTRIE" -> "Dupont Industrie"
    clients = {}
    for nom in liste_clients():
        clients[nettoyer(nom)] = nom

    a_creer = []
    problemes = []
    deja_vus = set()

    for numero_ligne, client, affaire, objet in lire_excel():
        affaire = affaire.upper().replace(" ", "")
        numero = affaire[:5]
        initiales = affaire[5:]
        objet = nettoyer(objet)

        if numero in existants:
            continue        # le dossier existe déjà : rien à faire

        if numero in deja_vus:
            problemes.append((numero_ligne, f"Le numéro {numero} apparaît plusieurs fois dans l'Excel"))
            continue
        deja_vus.add(numero)

        nom_client = clients.get(nettoyer(client))      # None si le client n'a pas de dossier
        if nom_client is None:
            problemes.append((numero_ligne, f"Client introuvable : {client}"))
            continue

        erreurs = verifier(nom_client, numero, initiales, objet)
        if erreurs:
            problemes.append((numero_ligne, " / ".join(erreurs)))
        else:
            a_creer.append((numero_ligne, nom_client, numero, initiales, objet))

    return a_creer, problemes


# ============================================================
# Interface graphique
# ============================================================

def filtrer_clients(event):
    """À chaque touche tapée dans le champ client, ne garde que les clients qui contiennent la saisie."""
    if event.keysym in ("Up", "Down", "Return", "Escape", "Tab"):
        return      # ces touches servent à naviguer, pas à écrire
    saisie = champ_client.get().upper().strip()
    resultats = []
    for client in tous_les_clients:
        if saisie in client.upper():
            resultats.append(client)
    champ_client["values"] = resultats


def activer_saut_clavier(combobox):
    """Liste ouverte : taper des lettres saute au premier client qui commence par ces lettres."""
    memoire = {"texte": "", "temps": 0, "liste": None}

    def sauter(lettre, temps):
        liste = memoire["liste"]
        temps = int(temps)
        if lettre == "" or not lettre.isprintable():
            return      # Entrée, flèches, Maj... : on laisse tkinter les gérer
        if temps - memoire["temps"] > 1000:     # plus d'une seconde depuis la dernière touche : on repart de zéro
            memoire["texte"] = ""
        memoire["texte"] += lettre.upper()
        memoire["temps"] = temps
        for i, client in enumerate(combobox["values"]):
            if client.upper().startswith(memoire["texte"]):
                combobox.tk.call(liste, "selection", "clear", 0, "end")
                combobox.tk.call(liste, "selection", "set", i)
                combobox.tk.call(liste, "activate", i)
                combobox.tk.call(liste, "see", i)
                break

    commande = combobox.register(sauter)

    def chercher_liste(widget):
        """Cherche la liste (Listbox) parmi les widgets internes de la Combobox, quel que soit son nom."""
        if combobox.tk.call("winfo", "class", widget) == "Listbox":
            return widget
        for enfant in combobox.tk.splitlist(combobox.tk.call("winfo", "children", widget)):
            trouve = chercher_liste(enfant)
            if trouve:
                return trouve
        return None

    def preparer():
        """Lancée à chaque ouverture de la liste : branche le saut clavier la première fois."""
        if memoire["liste"]:
            return      # déjà branché
        try:
            fenetre_liste = combobox.tk.call("ttk::combobox::PopdownWindow", combobox)
            liste = chercher_liste(fenetre_liste)
            if liste:
                combobox.tk.call("bind", liste, "<KeyPress>", f"+{commande} %A %t")
                memoire["liste"] = liste
        except tk.TclError:
            pass        # en cas d'échec, la liste marche quand même, juste sans saut clavier

    # postcommand : fonction appelée juste avant que la liste s'ouvre
    combobox.configure(postcommand=preparer)


def tableau(parent, colonnes, ligne):
    """Crée un tableau (Treeview) avec une barre de défilement. colonnes = [(titre, largeur), ...]"""
    cadre = ttk.Frame(parent)
    cadre.grid(row=ligne, column=0, sticky="nsew", padx=8)
    noms = [titre for titre, largeur in colonnes]
    arbre = ttk.Treeview(cadre, columns=noms, show="headings", height=8)
    for titre, largeur in colonnes:
        arbre.heading(titre, text=titre)
        arbre.column(titre, width=largeur, anchor="w")
    barre = ttk.Scrollbar(cadre, orient="vertical", command=arbre.yview)
    arbre.configure(yscrollcommand=barre.set)
    arbre.grid(row=0, column=0, sticky="nsew")
    barre.grid(row=0, column=1, sticky="ns")
    return arbre


def afficher_apercu(a_creer, problemes):
    """Ouvre une fenêtre avec les dossiers à créer et les lignes à corriger."""
    apercu = tk.Toplevel(fenetre)
    apercu.title("Comparaison avec l'Excel")
    apercu.transient(fenetre)       # garde l'aperçu au-dessus de la fenêtre principale

    ttk.Label(apercu, text=f"Dossiers à créer : {len(a_creer)}").grid(row=0, column=0, sticky="w", padx=8, pady=(8, 2))
    tableau_creer = tableau(apercu, [("Ligne", 50), ("Client", 200), ("Dossier", 320)], 1)
    for numero_ligne, client, numero, initiales, objet in a_creer:
        tableau_creer.insert("", tk.END, values=(numero_ligne, client, f"{numero}{initiales} {objet}"))

    ttk.Label(apercu, text=f"Lignes à corriger dans l'Excel : {len(problemes)}").grid(row=2, column=0, sticky="w", padx=8, pady=(10, 2))
    tableau_problemes = tableau(apercu, [("Ligne", 50), ("Problème", 520)], 3)
    for numero_ligne, probleme in problemes:
        tableau_problemes.insert("", tk.END, values=(numero_ligne, probleme))

    def creer_tout():
        crees = 0
        echecs = []
        for numero_ligne, client, numero, initiales, objet in a_creer:
            try:
                creer_affaire(client, numero, initiales, objet)
                crees += 1
            except OSError as erreur:
                echecs.append(f"Ligne {numero_ligne} : {erreur}")
        apercu.destroy()
        message = f"{crees} dossier(s) créé(s)."
        if echecs:
            messagebox.showwarning("Synchronisation terminée", message + "\n\nÉchecs :\n" + "\n".join(echecs))
        else:
            messagebox.showinfo("Synchronisation terminée", message)

    boutons = ttk.Frame(apercu)
    boutons.grid(row=4, column=0, sticky="e", padx=8, pady=10)
    bouton_creer = ttk.Button(boutons, text=f"Créer les {len(a_creer)} dossiers", command=creer_tout)
    bouton_creer.pack(side="left", padx=4)
    if not a_creer:
        bouton_creer.state(["disabled"])        # rien à créer : bouton grisé
    ttk.Button(boutons, text="Fermer", command=apercu.destroy).pack(side="left", padx=4)


def on_comparer():
    if EXCEL is None:
        messagebox.showerror("Excel non configuré", "Il manque la section « excel » dans config.json.")
        return

    bouton_excel.state(["disabled"])
    fenetre.config(cursor="watch")
    resultat = {}

    def travail():
        try:
            resultat["donnees"] = comparer_excel()
        except Exception as erreur:
            resultat["erreur"] = erreur

    thread = threading.Thread(target=travail, daemon=True)
    thread.start()
    fenetre.after(100, surveiller, thread, resultat)


def surveiller(thread, resultat):
    if thread.is_alive():
        fenetre.after(100, surveiller, thread, resultat)
        return

    bouton_excel.state(["!disabled"])
    fenetre.config(cursor="")

    erreur = resultat.get("erreur")
    if isinstance(erreur, FileNotFoundError):
        messagebox.showerror("Excel introuvable", f"Le fichier {DOSSIER_APP / EXCEL['chemin']} est introuvable.")
        return
    if isinstance(erreur, PermissionError):
        messagebox.showerror("Excel inaccessible", "Le fichier Excel est verrouillé. Réessaie dans un instant.")
        return
    if isinstance(erreur, KeyError):
        messagebox.showerror("Feuille introuvable", f"La feuille « {EXCEL['feuille']} » n'existe pas dans l'Excel.")
        return
    if erreur is not None:
        messagebox.showerror("Lecture impossible", f"Impossible de lire l'Excel :\n{erreur}")
        return

    a_creer, problemes = resultat["donnees"]
    if not a_creer and not problemes:
        messagebox.showinfo("Comparaison avec l'Excel", "Tout est à jour : aucun dossier à créer.")
    else:
        afficher_apercu(a_creer, problemes)


def on_creer():
    """Lancée au clic sur le bouton Créer."""
    global tous_les_clients
    client = nettoyer(champ_client.get())
    numero = champ_numero.get().strip()
    initiales = nettoyer(champ_initiales.get())
    objet = nettoyer(champ_objet.get())

    erreurs = verifier(client, numero, initiales, objet)
    if erreurs:
        messagebox.showerror("Saisie incorrecte", "\n".join(erreurs))
        return

    if not (BASE / client).is_dir():
        reponse = messagebox.askyesno("Nouveau client", f"Le client {client} n'existe pas.\n\nLe créer ?")
        if not reponse:
            return

    try:
        chemin = creer_affaire(client, numero, initiales, objet)
    except OSError as erreur:
        messagebox.showerror("Création impossible", str(erreur))
        return

    tous_les_clients = liste_clients()      # pour qu'un nouveau client apparaisse dans la liste
    champ_client["values"] = tous_les_clients
    champ_numero.delete(0, tk.END)
    champ_objet.delete(0, tk.END)

    if messagebox.askyesno("Dossier créé", f"{chemin.name} a été créé.\n\nOuvrir le dossier ?"):
        ouvrir_dossier(chemin)


fenetre = tk.Tk()
fenetre.title("Création de dossier d'affaire")

if not BASE.is_dir():
    messagebox.showerror("Dossier introuvable", f"Le dossier {BASE} est introuvable.")
    sys.exit()

tous_les_clients = liste_clients()      # la liste complète, lue une fois au démarrage

ttk.Label(fenetre, text="Client :").grid(row=0, column=0, sticky="w", padx=8, pady=5)
champ_client = ttk.Combobox(fenetre, values=tous_les_clients, width=40)
champ_client.grid(row=0, column=1, padx=8, pady=5)
champ_client.bind("<KeyRelease>", filtrer_clients)
activer_saut_clavier(champ_client)

ttk.Label(fenetre, text="N° d'affaire :").grid(row=1, column=0, sticky="w", padx=8, pady=5)
champ_numero = ttk.Entry(fenetre, width=43)
champ_numero.grid(row=1, column=1, padx=8, pady=5)

ttk.Label(fenetre, text="Initiales :").grid(row=2, column=0, sticky="w", padx=8, pady=5)
champ_initiales = ttk.Entry(fenetre, width=43)
champ_initiales.grid(row=2, column=1, padx=8, pady=5)

ttk.Label(fenetre, text="Objet :").grid(row=3, column=0, sticky="w", padx=8, pady=5)
champ_objet = ttk.Entry(fenetre, width=43)
champ_objet.grid(row=3, column=1, padx=8, pady=5)

# command=on_comparer SANS parenthèses : la comparaison ne se lance qu'au clic
bouton_excel = ttk.Button(fenetre, text="Comparer avec l'Excel", command=on_comparer)
bouton_excel.grid(row=4, column=0, sticky="w", padx=8, pady=10)

bouton = ttk.Button(fenetre, text="Créer", command=on_creer)
bouton.grid(row=4, column=1, sticky="e", padx=8, pady=10)

champ_client.focus()
fenetre.mainloop()